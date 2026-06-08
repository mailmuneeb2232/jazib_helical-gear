"""
Helical Gear Calculator — Flask Application
"""

import io
import json
import math
from flask import Flask, render_template, request, jsonify, send_file

from calculations.gear_calculator import calculate_gears, validate_inputs

app = Flask(__name__)
app.secret_key = "helical-gear-2024"

# --------------------------------------------------------------------------- #
#  Routes
# --------------------------------------------------------------------------- #

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/calculate", methods=["POST"])
def calculate():
    data = request.get_json(force=True)
    inputs_raw = data.get("inputs", {})
    selected_outputs = data.get("selected_outputs", [])

    if not selected_outputs:
        return jsonify({"error": "Please select at least one output parameter."}), 400

    parsed, errors = validate_inputs(inputs_raw)
    if errors:
        return jsonify({"errors": errors}), 422

    result = calculate_gears(parsed, selected_outputs)
    return jsonify(result)


@app.route("/export/pdf", methods=["POST"])
def export_pdf():
    """Generate and return a PDF calculation report."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch as rl_inch
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
        )
    except ImportError:
        return jsonify({"error": "ReportLab not installed. Run: pip install reportlab"}), 500

    data = request.get_json(force=True)
    inputs_raw = data.get("inputs", {})
    selected_outputs = data.get("selected_outputs", [])

    parsed, errors = validate_inputs(inputs_raw)
    if errors:
        return jsonify({"errors": errors}), 422

    calc = calculate_gears(parsed, selected_outputs)
    results = calc["results"]
    warnings = calc["warnings"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        leftMargin=0.75 * rl_inch, rightMargin=0.75 * rl_inch,
        topMargin=0.75 * rl_inch, bottomMargin=0.75 * rl_inch,
    )

    styles = getSampleStyleSheet()
    navy = colors.HexColor("#0d1b2a")
    blue = colors.HexColor("#1e90ff")

    title_style = ParagraphStyle(
        "Title", parent=styles["Title"],
        textColor=navy, fontSize=20, spaceAfter=4
    )
    sub_style = ParagraphStyle(
        "Sub", parent=styles["Normal"],
        textColor=blue, fontSize=10, spaceAfter=12
    )
    h2_style = ParagraphStyle(
        "H2", parent=styles["Heading2"],
        textColor=navy, fontSize=13, spaceBefore=12, spaceAfter=6
    )
    normal = styles["Normal"]

    story = []

    # Header
    story.append(Paragraph("Helical Gear Calculation Report", title_style))
    story.append(Paragraph("PD-Based Method · Inch Units · Engineering Grade", sub_style))
    story.append(HRFlowable(width="100%", thickness=1, color=blue))
    story.append(Spacer(1, 12))

    # Input parameters table
    story.append(Paragraph("Input Parameters", h2_style))
    input_labels = {
        "pitch_diameter": ("Pitch Diameter (PD)", "in"),
        "num_teeth": ("Number of Teeth (Z)", "—"),
        "helix_angle": ("Helix Angle (β)", "°"),
        "pressure_angle": ("Pressure Angle (φ)", "°"),
        "face_width": ("Face Width (F)", "in"),
        "power": ("Input Power (HP)", "HP"),
        "rpm": ("Rotational Speed", "RPM"),
        "pinion_teeth": ("Pinion Teeth (Z₁)", "—"),
        "gear_teeth": ("Gear Teeth (Z₂)", "—"),
        "service_factor": ("Service Factor (Ks)", "—"),
    }
    in_data = [["Parameter", "Value", "Unit"]]
    for key, (lbl, unit) in input_labels.items():
        val = parsed.get(key, "N/A")
        in_data.append([lbl, str(val), unit])

    in_table = Table(in_data, colWidths=[3.5 * rl_inch, 2 * rl_inch, 1.2 * rl_inch])
    in_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f0f4f8"), colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(in_table)
    story.append(Spacer(1, 16))

    # Results table
    story.append(Paragraph("Calculation Results", h2_style))
    res_data = [["Parameter", "Value", "Unit", "Formula"]]
    for key, r in results.items():
        res_data.append([r["label"], str(r["value"]), r["unit"], r["formula"]])

    res_table = Table(res_data, colWidths=[2.2 * rl_inch, 1.2 * rl_inch, 0.9 * rl_inch, 2.4 * rl_inch])
    res_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), blue),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#eaf4ff"), colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aaaaaa")),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("WORDWRAP", (3, 1), (3, -1), True),
    ]))
    story.append(res_table)

    # Warnings
    if warnings:
        story.append(Spacer(1, 14))
        story.append(Paragraph("Engineering Warnings", h2_style))
        for w in warnings:
            story.append(Paragraph(f"⚠  {w}", normal))
            story.append(Spacer(1, 4))

    # Footer
    story.append(Spacer(1, 20))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
    story.append(Paragraph(
        "Generated by Helical Gear Calculator · PD-Based Method · All dimensions in inches",
        ParagraphStyle("footer", parent=normal, fontSize=7, textColor=colors.grey)
    ))

    doc.build(story)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="helical_gear_report.pdf",
        mimetype="application/pdf"
    )


@app.route("/export/excel", methods=["POST"])
def export_excel():
    """Generate and return an Excel calculation report."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({"error": "openpyxl not installed. Run: pip install openpyxl"}), 500

    data = request.get_json(force=True)
    inputs_raw = data.get("inputs", {})
    selected_outputs = data.get("selected_outputs", [])

    parsed, errors = validate_inputs(inputs_raw)
    if errors:
        return jsonify({"errors": errors}), 422

    calc = calculate_gears(parsed, selected_outputs)
    results = calc["results"]
    warnings = calc["warnings"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gear Report"

    navy_fill   = PatternFill("solid", fgColor="0D1B2A")
    blue_fill   = PatternFill("solid", fgColor="1E90FF")
    light_fill  = PatternFill("solid", fgColor="EAF4FF")
    alt_fill    = PatternFill("solid", fgColor="F0F4F8")
    white_font  = Font(color="FFFFFF", bold=True)
    bold_font   = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    def hdr(ws, row, col, text, fill, font=None):
        c = ws.cell(row=row, column=col, value=text)
        c.fill = fill
        c.font = font or white_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    def cell(ws, row, col, text, fill=None):
        c = ws.cell(row=row, column=col, value=text)
        if fill:
            c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = thin_border

    # Title
    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = "Helical Gear Calculation Report — PD-Based Method (Inch Units)"
    t.font = Font(size=14, bold=True, color="0D1B2A")
    t.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28

    # Input parameters
    r = 3
    for c_idx, txt in enumerate(["Parameter", "Value", "Unit"], start=1):
        hdr(ws, r, c_idx, txt, navy_fill)
    r += 1

    input_labels = {
        "pitch_diameter": ("Pitch Diameter (PD)", "in"),
        "num_teeth": ("Number of Teeth (Z)", "—"),
        "helix_angle": ("Helix Angle (β)", "°"),
        "pressure_angle": ("Pressure Angle (φ)", "°"),
        "face_width": ("Face Width (F)", "in"),
        "power": ("Input Power (HP)", "HP"),
        "rpm": ("Rotational Speed", "RPM"),
        "pinion_teeth": ("Pinion Teeth (Z₁)", "—"),
        "gear_teeth": ("Gear Teeth (Z₂)", "—"),
        "service_factor": ("Service Factor (Ks)", "—"),
    }
    for i, (key, (lbl, unit)) in enumerate(input_labels.items()):
        fill = alt_fill if i % 2 == 0 else None
        cell(ws, r, 1, lbl, fill)
        cell(ws, r, 2, parsed.get(key, "N/A"), fill)
        cell(ws, r, 3, unit, fill)
        r += 1

    # Results
    r += 1
    for c_idx, txt in enumerate(["Parameter", "Value", "Unit", "Formula", "Description"], start=1):
        hdr(ws, r, c_idx, txt, blue_fill)
    r += 1
    for i, (key, res) in enumerate(results.items()):
        fill = light_fill if i % 2 == 0 else None
        cell(ws, r, 1, res["label"], fill)
        cell(ws, r, 2, res["value"], fill)
        cell(ws, r, 3, res["unit"], fill)
        cell(ws, r, 4, res["formula"], fill)
        cell(ws, r, 5, res["description"], fill)
        r += 1

    # Warnings
    if warnings:
        r += 1
        ws.cell(row=r, column=1, value="⚠  Engineering Warnings").font = bold_font
        r += 1
        for w in warnings:
            ws.cell(row=r, column=1, value=w)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            r += 1

    # Column widths
    for col, width in zip("ABCDE", [32, 14, 10, 45, 50]):
        ws.column_dimensions[get_column_letter(ord(col) - 64)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="helical_gear_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
