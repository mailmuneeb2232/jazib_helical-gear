"""
Engineering Report Generator for Helical Gear Design
Generates PDF and Excel reports
"""
import datetime


class ReportGenerator:
    def __init__(self):
        self.timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def generate_text_report(self, inputs, outputs, assumptions, calc_details, design_status):
        """Generate a plain text engineering report"""
        lines = []
        lines.append("=" * 70)
        lines.append("       HELICAL GEAR DESIGN REPORT")
        lines.append("       AGMA Standards - Imperial Units (Inch System)")
        lines.append(f"       Generated: {self.timestamp}")
        lines.append("=" * 70)
        lines.append("")

        lines.append("─" * 70)
        lines.append("1. INPUT PARAMETERS")
        lines.append("─" * 70)
        for key, (label, value, unit) in inputs.items():
            lines.append(f"   {label:<45} {str(value):>12} {unit}")

        lines.append("")
        lines.append("─" * 70)
        lines.append("2. ASSUMED VALUES (AGMA Standard or User-Specified)")
        lines.append("─" * 70)
        if assumptions:
            for key, (label, value, unit, source) in assumptions.items():
                lines.append(f"   {label:<45} {str(value):>12} {unit}   [{source}]")
        else:
            lines.append("   No assumptions made.")

        lines.append("")
        lines.append("─" * 70)
        lines.append("3. CALCULATION DETAILS")
        lines.append("─" * 70)
        for key, detail in calc_details.items():
            lines.append(f"\n   [{detail['label']}]")
            lines.append(f"   Formula    : {detail['formula']}")
            lines.append(f"   Substitution: {detail['substitution']}")
            lines.append(f"   Result     : {detail['result']:.6g} {detail['unit']}")

        lines.append("")
        lines.append("─" * 70)
        lines.append("4. FINAL RESULTS")
        lines.append("─" * 70)
        for key, detail in calc_details.items():
            lines.append(f"   {detail['label']:<45} {detail['result']:>14.4g} {detail['unit']}")

        lines.append("")
        lines.append("─" * 70)
        lines.append("5. DESIGN VERIFICATION")
        lines.append("─" * 70)
        if design_status:
            lines.append(f"   STATUS: {design_status['status']}")
            for note in design_status.get('notes', []):
                lines.append(f"   • {note}")
        else:
            lines.append("   No stress analysis performed.")

        lines.append("")
        lines.append("=" * 70)
        lines.append("   END OF REPORT")
        lines.append("=" * 70)

        return "\n".join(lines)

    def export_pdf(self, report_text, filename):
        """Export report as PDF using reportlab if available, else HTML"""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER, TA_LEFT

            doc = SimpleDocTemplate(filename, pagesize=letter,
                                    rightMargin=0.75 * inch, leftMargin=0.75 * inch,
                                    topMargin=0.75 * inch, bottomMargin=0.75 * inch)
            styles = getSampleStyleSheet()
            story = []

            title_style = ParagraphStyle('title', parent=styles['Title'],
                                          fontSize=16, textColor=colors.HexColor('#1e66f5'),
                                          spaceAfter=12, alignment=TA_CENTER)
            heading_style = ParagraphStyle('heading', parent=styles['Heading2'],
                                            fontSize=12, textColor=colors.HexColor('#1e66f5'),
                                            spaceBefore=12, spaceAfter=6)
            body_style = ParagraphStyle('body', parent=styles['Normal'],
                                         fontSize=9, fontName='Courier',
                                         leading=14)

            story.append(Paragraph("HELICAL GEAR DESIGN REPORT", title_style))
            story.append(Paragraph("AGMA Standards – Imperial Units (Inch System)", styles['Normal']))
            story.append(Paragraph(f"Generated: {self.timestamp}", styles['Normal']))
            story.append(Spacer(1, 12))

            for line in report_text.split('\n'):
                if line.startswith('=') or line.startswith('─'):
                    continue
                elif line.strip().startswith(('1.', '2.', '3.', '4.', '5.')):
                    story.append(Paragraph(line.strip(), heading_style))
                else:
                    safe_line = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(safe_line.replace(' ', '&nbsp;'), body_style))

            doc.build(story)
            return True, ""
        except ImportError:
            html = f"<html><body><pre style='font-family:monospace'>{report_text}</pre></body></html>"
            html_file = filename.replace('.pdf', '.html')
            with open(html_file, 'w') as f:
                f.write(html)
            return False, f"reportlab not installed. Saved as HTML: {html_file}"

    def export_excel(self, inputs, outputs, assumptions, calc_details, filename):
        """Export report as Excel file using openpyxl if available"""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Gear Design Report"

            header_font = Font(bold=True, size=12, color="FFFFFF")
            header_fill = PatternFill(fill_type="solid", fgColor="1E66F5")
            section_font = Font(bold=True, size=11, color="1E66F5")
            section_fill = PatternFill(fill_type="solid", fgColor="E8F0FE")
            normal_font = Font(size=10)

            row = 1
            ws.cell(row, 1, "HELICAL GEAR DESIGN REPORT").font = Font(bold=True, size=14, color="1E66F5")
            ws.merge_cells(f'A{row}:E{row}')
            row += 1
            ws.cell(row, 1, f"Generated: {self.timestamp}")
            row += 2

            def write_section(title, data_dict, columns):
                nonlocal row
                ws.cell(row, 1, title).font = section_font
                ws.cell(row, 1).fill = section_fill
                ws.merge_cells(f'A{row}:E{row}')
                row += 1
                for c, header in enumerate(columns, 1):
                    cell = ws.cell(row, c, header)
                    cell.font = header_font
                    cell.fill = header_fill
                row += 1
                for key, vals in data_dict.items():
                    for c, val in enumerate(vals, 1):
                        ws.cell(row, c, val).font = normal_font
                    row += 1
                row += 1

            write_section("INPUT PARAMETERS",
                          {k: [v[0], str(v[1]), v[2]] for k, v in inputs.items()},
                          ["Parameter", "Value", "Unit"])

            if assumptions:
                write_section("ASSUMED VALUES",
                              {k: [v[0], str(v[1]), v[2], v[3]] for k, v in assumptions.items()},
                              ["Parameter", "Value", "Unit", "Source"])

            write_section("CALCULATED RESULTS",
                          {k: [d['label'], f"{d['result']:.6g}", d['unit'], d['formula']]
                           for k, d in calc_details.items()},
                          ["Parameter", "Result", "Unit", "Formula"])

            for col in range(1, 6):
                ws.column_dimensions[get_column_letter(col)].width = 30

            wb.save(filename)
            return True, ""
        except ImportError:
            import csv
            csv_file = filename.replace('.xlsx', '.csv')
            with open(csv_file, 'w', newline='') as f:
                writer = csv.writer(f)
                writer.writerow(["Parameter", "Value", "Unit"])
                for key, vals in inputs.items():
                    writer.writerow([vals[0], str(vals[1]), vals[2]])
                writer.writerow([])
                writer.writerow(["Calculated Results"])
                for key, d in calc_details.items():
                    writer.writerow([d['label'], f"{d['result']:.6g}", d['unit']])
            return False, f"openpyxl not installed. Saved as CSV: {csv_file}"
